# -*- coding: utf-8 -*-
"""
coletar_dados.py
================
Extrai os dados da planilha-mestra DataBase/Dados_GeraisReport_AFR.xlsx
e gera os JSONs consumidos pelo site em /dados:

    dados/fisico_total.json   <- aba "% Físico Total"
    dados/pluviometria.json   <- aba "Pluviometria"
    dados/status_semanal.json <- aba "Status Semanal"
    dados/versao.json         <- metadados (data de geração)

Uso:
    python gerador/coletar_dados.py

A planilha é a FONTE DE ATUALIZAÇÃO SEMANAL. Sempre que ela for
atualizada, basta rodar este script para regenerar os JSONs.
"""
from __future__ import annotations

import json
import os
import re
import sys
import unicodedata
from datetime import datetime, date

import openpyxl

# --------------------------------------------------------------------
# Caminhos
# --------------------------------------------------------------------
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(RAIZ, 'DataBase', 'Dados_GeraisReport_AFR.xlsx')
SAIDA = os.path.join(RAIZ, 'dados')


# --------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------
def normalizar(txt) -> str:
    """Normaliza um rótulo para comparação (sem acento, minúsculo, espaços)."""
    if txt is None:
        return ''
    s = str(txt).strip()
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', s).lower()


def achar_aba(wb, alvo: str):
    """Localiza uma aba pelo nome normalizado (tolerante a encoding)."""
    alvo_n = normalizar(alvo)
    for nome in wb.sheetnames:
        if normalizar(nome) == alvo_n:
            return nome
    for nome in wb.sheetnames:
        if alvo_n in normalizar(nome):
            return nome
    raise KeyError(f'Aba não encontrada: {alvo}')


def achar_coluna(cabecalhos, *apelidos):
    """Índice da 1ª coluna cujo cabeçalho normalizado casa com algum apelido."""
    norm = [normalizar(h) for h in cabecalhos]
    for apelido in apelidos:
        ap = normalizar(apelido)
        for i, h in enumerate(norm):
            if h == ap:
                return i
    for apelido in apelidos:
        ap = normalizar(apelido)
        for i, h in enumerate(norm):
            if ap and ap in h:
                return i
    raise KeyError(f'Coluna não encontrada: {apelidos} em {cabecalhos}')


def iso_data(v):
    """Converte datetime/date/str para 'YYYY-MM-DD' (ou None)."""
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    if not s:
        return None
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return f'{m.group(1)}-{m.group(2)}-{m.group(3)}'
    return s


def num(v):
    """Converte para float (ou None)."""
    if v is None or v == '':
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def inteiro(v):
    if v is None or v == '':
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------
# Aba: % Físico Total  ->  fisico_total.json
# --------------------------------------------------------------------
def ler_fisico(wb):
    ws = wb[achar_aba(wb, '% Físico Total')]
    linhas = ws.iter_rows(values_only=True)
    cab = list(next(linhas))

    c_id     = achar_coluna(cab, 'ID Contrato', 'id_contrato')
    c_contr  = achar_coluna(cab, 'Contratada')
    c_mes    = achar_coluna(cab, 'Mês', 'Mes')
    c_semana = achar_coluna(cab, 'Semana')
    c_data   = achar_coluna(cab, 'Data')
    c_b1     = achar_coluna(cab, 'Baseline 1')          # exato (não pega "Acum")
    c_b1ac   = achar_coluna(cab, 'Baseline 1 Acum')
    c_real   = achar_coluna(cab, 'Real')
    c_realac = achar_coluna(cab, 'Real Acum')
    c_ader   = achar_coluna(cab, 'Aderência', 'Aderencia')
    # Projetado Acum (série de tendência da curva S)
    try:
        c_projac = achar_coluna(cab, 'Projetado Acum')
    except KeyError:
        c_projac = None

    out = []
    for r in linhas:
        if r is None or all(x is None for x in r):
            continue
        if r[c_id] is None and r[c_data] is None:
            continue
        reg = {
            'id_contrato': str(r[c_id]).strip() if r[c_id] is not None else None,
            'contratada': str(r[c_contr]).strip() if r[c_contr] is not None else None,
            'mes': iso_data(r[c_mes]),
            'semana': inteiro(r[c_semana]),
            'data': iso_data(r[c_data]),
            'baseline_1_acum': num(r[c_b1ac]),
            'real_acum': num(r[c_realac]),
            'aderencia': num(r[c_ader]),
            'baseline_1': num(r[c_b1]),
            'real': num(r[c_real]),
        }
        if c_projac is not None:
            reg['projetado_acum'] = num(r[c_projac])
        out.append(reg)
    return out


# --------------------------------------------------------------------
# Aba: Pluviometria  ->  pluviometria.json
# --------------------------------------------------------------------
def ler_pluviometria(wb):
    ws = wb[achar_aba(wb, 'Pluviometria')]
    linhas = ws.iter_rows(values_only=True)
    cab = list(next(linhas))

    c_mes   = achar_coluna(cab, 'Mês', 'Mes')
    c_sem   = achar_coluna(cab, 'Semana')
    c_data  = achar_coluna(cab, 'Data')
    c_histA = achar_coluna(cab, 'histórica acum', 'historica acum')
    c_histM = achar_coluna(cab, 'histórica média/dia', 'historica media/dia',
                           'histórica média', 'historica media')
    c_mmdia = achar_coluna(cab, 'mm/dia', 'mm dia')
    c_mmac  = achar_coluna(cab, 'mm acum')
    try:
        c_par = achar_coluna(cab, 'dia paralisado?', 'dia paralisado')
    except KeyError:
        c_par = None

    out = []
    for r in linhas:
        if r is None or all(x is None for x in r):
            continue
        if r[c_data] is None:
            continue
        par = r[c_par] if c_par is not None else None
        par_norm = None
        if par is not None and str(par).strip() != '':
            par_norm = str(par).strip().lower()
        out.append({
            'mes': iso_data(r[c_mes]),
            'semana': inteiro(r[c_sem]),
            'data': iso_data(r[c_data]),
            'historica_acum': num(r[c_histA]),
            'historica_media_dia': num(r[c_histM]),
            'mm_dia': num(r[c_mmdia]),
            'mm_acum': num(r[c_mmac]),
            'dia_paralisado': par_norm,
        })
    return out


# --------------------------------------------------------------------
# Aba: Status Semanal  ->  status_semanal.json
# --------------------------------------------------------------------
def ler_status(wb):
    ws = wb[achar_aba(wb, 'Status Semanal')]
    linhas = ws.iter_rows(values_only=True)
    cab = list(next(linhas))

    c_contr = achar_coluna(cab, 'Contratada')
    c_mesref = achar_coluna(cab, 'Mês Referência', 'Mes Referencia', 'Mês')
    c_sem   = achar_coluna(cab, 'Semana')
    c_item  = achar_coluna(cab, 'Item')
    c_desc  = achar_coluna(cab, 'Descrição', 'Descricao')

    out = []
    for r in linhas:
        if r is None or all(x is None for x in r):
            continue
        desc = r[c_desc]
        if desc is None:
            continue
        out.append({
            'contratada': str(r[c_contr]).strip() if r[c_contr] is not None else None,
            'mes_referencia': iso_data(r[c_mesref]),
            'semana': inteiro(r[c_sem]),
            'item': str(r[c_item]).strip() if r[c_item] is not None else None,
            'descricao': str(desc).strip(),
        })
    return out


# --------------------------------------------------------------------
# main
# --------------------------------------------------------------------
def escrever(nome, dados):
    caminho = os.path.join(SAIDA, nome)
    with open(caminho, 'w', encoding='utf-8') as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)
    print(f'  {nome}: {len(dados)} registros')


def main():
    if not os.path.exists(XLSX):
        print(f'ERRO: planilha não encontrada em {XLSX}', file=sys.stderr)
        return 1
    os.makedirs(SAIDA, exist_ok=True)

    print(f'Lendo {XLSX} ...')
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    try:
        fisico = ler_fisico(wb)
        pluvio = ler_pluviometria(wb)
        status = ler_status(wb)
    finally:
        wb.close()

    print('Gravando JSONs em /dados:')
    escrever('fisico_total.json', fisico)
    escrever('pluviometria.json', pluvio)
    escrever('status_semanal.json', status)

    versao = {
        'gerado_em': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'fonte': os.path.basename(XLSX),
        'registros': {
            'fisico_total': len(fisico),
            'pluviometria': len(pluvio),
            'status_semanal': len(status),
        },
    }
    escrever('versao.json', versao)
    print('OK.')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
