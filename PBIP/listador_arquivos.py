from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def gerar_excel_arquivos(diretorio_inicial="."):
    pasta_raiz = Path(diretorio_inicial).resolve()
    print(f"\nA pesquisar em: {pasta_raiz}\n")
    
    # Criar workbook do Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Arquivos Listados"
    
    # Ativar linhas de grade no Excel
    ws.views.sheetView[0].showGridLines = True
    
    # Cabeçalhos das colunas
    headers = ["Local (Raiz)", "Arquivo (Nome)", "Tipo (Extensão)"]
    ws.append(headers)
    
    # Estilizar o cabeçalho (Azul escuro com texto branco)
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        
    ws.row_dimensions[1].height = 25
    
    # Percorrer diretórios recursivamente (pastas e subpastas)
    contador = 0
    row_idx = 2
    for caminho in pasta_raiz.rglob("*"):
        if caminho.is_file():
            local = str(caminho.parent)
            nome_ficheiro = caminho.name
            extensao = caminho.suffix.lstrip('.') or "sem_extensao"
            
            # Adicionar linha no Excel
            ws.append([local, nome_ficheiro, extensao])
            
            # Estilizar células de dados
            for col_num in range(1, 4):
                c = ws.cell(row=row_idx, column=col_num)
                c.font = Font(name="Segoe UI", size=10)
                c.border = thin_border
                c.alignment = Alignment(vertical="center")
                
            ws.row_dimensions[row_idx].height = 20
            row_idx += 1
            contador += 1
            
    # Ajustar largura das colunas automaticamente para melhor legibilidade
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = max(max_length + 4, 15)
        ws.column_dimensions[column].width = min(adjusted_width, 60)
        
    # Guardar o ficheiro Excel
    nome_arquivo_saida = "inventario_arquivos.xlsx"
    wb.save(nome_arquivo_saida)
    print(f"\nConcluído! {contador} ficheiros encontrados.")
    print(f"Ficheiro Excel gerado com sucesso: {nome_arquivo_saida}\n")

if __name__ == "__main__":
    gerar_excel_arquivos()
